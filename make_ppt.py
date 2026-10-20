"""Builds the SpamSentry project presentation (9 slides).
ACC is filled in from the trained weights.json so the results slide
shows the real 5-fold CV number.
"""
import json

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls


def set_cell_borders(cell, color="000000", w="12700"):  # w in EMU; 12700 = 1pt
    """Give a table cell plain black grid borders via raw XML
    (python-pptx exposes no public border API)."""
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnT", "a:lnL", "a:lnB", "a:lnR"):
        ln = parse_xml(
            f'<{tag} {nsdecls("a")} w="{w}">'
            f'<a:solidFill><a:srgbClr val="{color}"/></a:solidFill></{tag}>'
        )
        tcPr.append(ln)

with open("weights.json") as f:
    W = json.load(f)
ACC = W["accuracy"]
ACC_PCT = f"{ACC * 100:.1f}%"
NFEAT = len(W["vocab"])
N_SPAM = W.get("n_spam", 108)
N_HAM = W.get("n_ham", 108)
FOLDS = W.get("fold_accuracies", [])

BLACK = RGBColor(0x00, 0x00, 0x00)
GREY = RGBColor(0x44, 0x44, 0x44)
FONT = "Times New Roman"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def txbox(slide, left, top, width, height, text, size=20, bold=False,
          color=BLACK, align=None, font=FONT):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    p = tb.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font
    if align is not None:
        p.alignment = align
    return tb


def number_tag(slide, n):
    txbox(slide, 0.8, 0.35, 2, 0.6, f"{n:02d}", size=18, bold=True, color=GREY)


def content_slide(title, n, body_lines, table_rows=None):
    """Title + body text lines, optional plain black grid table under the text."""
    s = prs.slides.add_slide(blank)
    number_tag(s, n)
    txbox(s, 0.8, 0.85, 11.7, 1.1, title, size=36, bold=True)
    y = 2.05
    for line in body_lines:
        txbox(s, 0.8, y, 11.7, 0.6, line, size=20)
        y += 0.55
    if table_rows:
        rows, cols = len(table_rows), len(table_rows[0])
        g = s.shapes.add_table(rows, cols, Inches(0.8), Inches(y + 0.25),
                               Inches(8.5), Inches(0.45 * rows))
        tbl = g.table
        for r in range(rows):
            for c in range(cols):
                cell = tbl.cell(r, c)
                cell.text = table_rows[r][c]
                for p in cell.text_frame.paragraphs:
                    p.font.size = Pt(16)
                    p.font.name = FONT
                    p.font.color.rgb = BLACK
                    p.font.bold = (r == 0)
                # plain black grid borders
                set_cell_borders(cell, "000000", "12700")
        for r in range(rows):
            tbl.rows[r].height = Inches(0.45)
    return s


# 1 — Title
s = prs.slides.add_slide(blank)
txbox(s, 1.0, 1.6, 11.3, 1.4, "SpamSentry", size=56, bold=True, align=PP_ALIGN.CENTER)
txbox(s, 1.0, 3.1, 11.3, 1.0, "AI-Powered SMS Spam Detector", size=28,
      align=PP_ALIGN.CENTER)
txbox(s, 1.0, 4.4, 11.3, 2.0,
      "Vikranth Reddy  |  B.Tech CSE (AI & ML), Woxsen University\n"
      "Daily AI/ML Project Factory — Day 3  ·  20 October 2026",
      size=20, color=GREY, align=PP_ALIGN.CENTER)

# 2 — Problem
content_slide("The Problem", 2, [
    "•  Spam texts waste our time and trick people into scams — free prizes,",
    "    fake lotteries, phishing links. Billions are sent every day.",
    "•  Phones get hundreds of messages; no one can read and judge each one.",
    "•  Rule-based filters break: spammers change wording, use caps and symbols.",
    "•  Goal: paste a message → get an instant SPAM vs HAM verdict from a real ML model.",
])

# 3 — Solution
content_slide("The Solution", 3, [
    "•  SpamSentry: an SMS spam detector trained on 216 labeled messages (108 spam, 108 ham).",
    "•  Pipeline: text → TF-IDF word features → Multinomial Naive Bayes → spam probability.",
    "•  Trained weights (vocab, idf, priors, log-probabilities) are exported to weights.json —",
    "    the website re-runs the exact same math in JavaScript. Zero backend.",
    "•  Verified: browser predictions match sklearn's predict_proba exactly.",
])

# 4 — Tech Stack
content_slide("Tech Stack", 4, [
    "•  Python, scikit-learn, NumPy — training & evaluation (train.py).",
    "•  TF-IDF vectorizer: unigrams, English stop-words removed, sublinear TF,",
    "    L2 normalization, top 200 features.",
    "•  Multinomial Naive Bayes classifier — ideal for word-count text data.",
    "•  Vanilla JavaScript — in-browser inference, no frameworks, no API calls.",
    "•  Static hosting on Vercel with hardened security headers.",
])

# 5 — Architecture
content_slide("Architecture", 5, [
    "•  train.py: (1) generate 216 SMS from spam/ham templates (seed 42) →",
    "    (2) stratified 5-fold CV → (3) TF-IDF fit → (4) Naive Bayes fit → (5) export weights.json.",
    "•  weights.json: 200-word vocabulary, idf values, class priors, feature log-probabilities,",
    "    stop-word list — tiny enough to ship to the browser.",
    "•  index.html: tokenizer (sklearn-compatible) → TF-IDF → per-class joint log-likelihood →",
    "    softmax → SPAM/HAM verdict + probability bar + top contributing words.",
    "•  Same input in Python and JS gives the same probability — parity verified.",
])

# 6 — Demo
content_slide("Live Demo", 6, [
    "•  Type or paste any message, hit Check → SPAM or HAM verdict with a probability bar.",
    "•  Try a spam example: \"WINNER!! You won a free iPhone, text CLAIM to 80808 now\" → SPAM.",
    "•  Try a ham example: \"Hey, are we still meeting for dinner at 7pm?\" → HAM.",
    "•  The demo also highlights which words pushed the decision (e.g. winner, claim vs dinner).",
    "•  Demo URL: spamsentry.vercel.app",
])

# 7 — Results (with plain black grid table of fold scores)
fold_pct = [f"{a * 100:.1f}" for a in FOLDS] if FOLDS else ["—"] * 5
fold_rows = [["Fold", "1", "2", "3", "4", "5", "Mean"],
             ["Accuracy (%)"] + fold_pct + [f"{ACC * 100:.1f}"]]
content_slide("Results", 7, [
    f"•  5-fold cross-validation accuracy: {ACC_PCT} ({N_SPAM + N_HAM} SMS: {N_SPAM} spam, {N_HAM} ham, stratified folds).",
    f"•  Model: TF-IDF ({NFEAT} unigram features) + Multinomial Naive Bayes.",
    "•  Most spammy words learned: winner, claim, free, prize, congratulations.",
    "•  Known limitation: brand-new spam tricks not in the templates may slip through —",
    "    retraining on real reported spam fixes it.",
], table_rows=fold_rows)

# 8 — Future Scope
content_slide("Future Scope", 8, [
    "•  Train on a real public SMS dataset (e.g. the UCI SMS Spam Collection, 5,574 messages).",
    "•  Add bigrams (\"free prize\", \"text claim\") to catch multi-word spam phrases.",
    "•  Compare models: logistic regression, linear SVM, then a small transformer.",
    "•  Add a browser extension / Android app that scans incoming messages live.",
    "•  Multilingual spam detection (Telugu + English code-mixed messages).",
])

# 9 — Thank You
s = prs.slides.add_slide(blank)
txbox(s, 1.5, 2.2, 10.3, 1.6, "Thank You", size=60, bold=True, align=PP_ALIGN.CENTER)
txbox(s, 1.5, 4.3, 10.3, 1.5,
      "GitHub: github.com/HVRGAMING2705/spamsentry\nDemo: spamsentry.vercel.app",
      size=22, color=GREY, align=PP_ALIGN.CENTER)

out = "/home/hatch/workspace/goals/daily-ai-ml-project-factory/files/AIML_SpamSentry_2026-10-20.pptx"
prs.save(out)
print("saved", out)
print(f"slides: {len(prs.slides)}, CV accuracy baked in: {ACC_PCT}")
